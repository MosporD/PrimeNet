"""
NE Comparison Routes
Handles XML configuration comparison functionality
"""

from flask import Blueprint, request, jsonify, send_file, render_template, redirect, url_for
from werkzeug.utils import secure_filename
import os
import tempfile
import uuid
from functools import wraps

from ncm_core import XMLComparator
from database_enhanced import get_user_by_session, log_activity

ne_comparison_bp = Blueprint(
    'ne_comparison', __name__,
    template_folder='templates',
    static_folder='static',
    static_url_path='/ne_comparison/static',
)

TEMP_FILES = {}

def login_required(f):
    """Decorator to require login"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        session_token = request.cookies.get('session_token')
        if not session_token:
            return redirect(url_for('auth.login_page'))

        user = get_user_by_session(session_token)
        if not user:
            return redirect(url_for('auth.login_page'))

        request.current_user = user
        return f(*args, **kwargs)

    return decorated_function

def get_current_user():
    """Get current logged-in user"""
    session_token = request.cookies.get('session_token')
    if session_token:
        return get_user_by_session(session_token)
    return None
def format_user_data(user):
    """Format user data for templates"""
    if not user:
        return None
    if isinstance(user, dict):
        return {'username': user.get('username'), 'email': user.get('email'), 'role': user.get('role'), 'id': user.get('id')}
    return {'username': (user.get('username') if isinstance(user, dict) else user[1]), 'email': (user.get('email') if isinstance(user, dict) else user[2]), 'role': (user.get('role') if isinstance(user, dict) else user[6]), 'id': (user.get('id') if isinstance(user, dict) else user[0])}


@ne_comparison_bp.route('/ne-comparison')
@login_required
def ne_comparison_page():
    """Render NE Comparison page"""
    user = get_current_user()
    return render_template('ne_comparison.html', user=format_user_data(user))

@ne_comparison_bp.route('/api/ne-comparison/compare', methods=['POST'])
def compare_files():
    """Compare two XML files - returns Excel file directly like old version"""
    user = get_current_user()
    if not user:
        return jsonify({'error': 'Unauthorized'}), 401

    if 'file1' not in request.files or 'file2' not in request.files:
        return jsonify({'error': 'Both files required'}), 400

    file1 = request.files['file1']
    file2 = request.files['file2']

    if file1.filename == '' or file2.filename == '':
        return jsonify({'error': 'Both files must be selected'}), 400

    if not (file1.filename.endswith('.xml') and file2.filename.endswith('.xml')):
        return jsonify({'error': 'Both files must be XML'}), 400

    try:
        file_id = str(uuid.uuid4())

        filename1 = secure_filename(file1.filename)
        filename2 = secure_filename(file2.filename)

        temp_path1 = os.path.join(tempfile.gettempdir(), f"{file_id}_1_{filename1}")
        temp_path2 = os.path.join(tempfile.gettempdir(), f"{file_id}_2_{filename2}")

        file1.save(temp_path1)
        file2.save(temp_path2)

        # Create output filename with timestamp like old version
        from datetime import datetime
        output_filename = f"XML_Comparison_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output_path = os.path.join(tempfile.gettempdir(), f"{file_id}_{output_filename}")

        # Compare
        comparator = XMLComparator(temp_path1, temp_path2, output_path)
        success, diff_count = comparator.compare()

        if not success:
            return jsonify({'error': 'Comparison failed'}), 500

        # Log activity
        log_activity((user.get('id') if isinstance(user, dict) else user[0]), 'ne_comparison', f'Compared {filename1} and {filename2}')

        # Send file directly like old version
        return send_file(
            output_path,
            as_attachment=True,
            download_name=output_filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@ne_comparison_bp.route('/api/ne-comparison/download-report', methods=['POST'])
def download_report():
    """Download comparison report"""
    user = get_current_user()
    if not user:
        return jsonify({'error': 'Unauthorized'}), 401

    try:
        data = request.get_json()

        import openpyxl
        from openpyxl.styles import Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Comparison Report"

        ws['A1'] = 'Type'
        ws['B1'] = 'Parameter/MO'
        ws['C1'] = 'Old Value'
        ws['D1'] = 'New Value'
        ws['E1'] = 'Path'

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")

        row = 2
        for diff in data.get('differences', []):
            ws[f'A{row}'] = diff.get('type', '')
            ws[f'B{row}'] = diff.get('parameter', diff.get('mo_class', ''))
            ws[f'C{row}'] = str(diff.get('old_value', ''))
            ws[f'D{row}'] = str(diff.get('new_value', ''))
            ws[f'E{row}'] = diff.get('path', '')

            if diff.get('type') == 'added':
                fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
            elif diff.get('type') == 'removed':
                fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
            elif diff.get('type') == 'modified':
                fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
            else:
                fill = None

            if fill:
                for cell in ws[row]:
                    cell.fill = fill

            row += 1

        temp_path = os.path.join(tempfile.gettempdir(), f"report_{uuid.uuid4()}.xlsx")
        wb.save(temp_path)

        log_activity((user.get('id') if isinstance(user, dict) else user[0]), 'report_download', 'Downloaded comparison report')

        return send_file(
            temp_path,
            as_attachment=True,
            download_name='comparison_report.xlsx'
        )

    except Exception as e:
        return jsonify({'error': str(e)}), 500
