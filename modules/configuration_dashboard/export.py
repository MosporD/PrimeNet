"""Excel export for Configuration Dashboard WNCELG site summary."""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _safe_filename_part(value: str, *, fallback: str = 'export') -> str:
    cleaned = re.sub(r'[^\w\-]+', '_', (value or '').strip())
    return cleaned.strip('_')[:40] or fallback


EXPORT_COLUMNS: tuple[tuple[str, str], ...] = (
    ('area', 'Area'),
    ('site_id', 'Site ID'),
    ('site_name', 'Site Name'),
    ('status', 'Status'),
    ('group_count', 'WNCELG Count'),
    ('instances_csv', 'WNCELG Instances'),
    ('dns_csv', 'DNs'),
    ('metadata_site_id', 'Metadata Site ID'),
)


def build_wncelg_workbook(payload: dict[str, Any]) -> tuple[io.BytesIO, str]:
    sites = payload.get('sites') or []
    if not sites:
        raise ValueError('No sites to export')

    area = str(payload.get('area') or 'all').strip() or 'all'
    status = str(payload.get('status') or 'all').strip() or 'all'
    username = str(payload.get('username') or '').strip()
    built_at = str(payload.get('built_at') or '').strip()
    mo_class = str(payload.get('mo_class') or '').strip()
    summary = payload.get('summary') or {}

    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    area_part = _safe_filename_part(area, fallback='all')
    filename = f'WNCELG_Site_Split_{area_part}_{stamp}.xlsx'

    wb = Workbook()
    hdr_fill = PatternFill(start_color='1F6FEB', end_color='1F6FEB', fill_type='solid')
    hdr_font = Font(color='FFFFFF', bold=True)

    ws_meta = wb.active
    ws_meta.title = 'Report Info'
    ws_meta.column_dimensions['A'].width = 28
    ws_meta.column_dimensions['B'].width = 56
    generated = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    meta_rows = [
        ('Report', 'Configuration Dashboard — WNCELG'),
        ('Generated (UTC)', generated),
        ('Exported By', username),
        ('Snapshot Built At', built_at or '(unknown)'),
        ('Area Filter', area),
        ('Status Filter', status),
        ('MO Class', mo_class),
        ('Total Sites', summary.get('total_sites', len(sites))),
        ('Split', summary.get('split', '')),
        ('No Split', summary.get('no_split', '')),
        ('Exported Rows', len(sites)),
    ]
    for col, label in enumerate(('Field', 'Value'), 1):
        cell = ws_meta.cell(row=1, column=col, value=label)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal='center')
    for row_idx, (field, value) in enumerate(meta_rows, start=2):
        ws_meta.cell(row=row_idx, column=1, value=field)
        ws_meta.cell(row=row_idx, column=2, value=value)

    ws = wb.create_sheet('WNCELG Sites')
    for col_idx, (_key, header) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = hdr_fill
        cell.font = hdr_font

    for row_idx, site in enumerate(sites, start=2):
        for col_idx, (key, _header) in enumerate(EXPORT_COLUMNS, start=1):
            value = site.get(key)
            if key == 'instances_csv' and not value and isinstance(site.get('instances'), list):
                value = ','.join(str(v) for v in site['instances'])
            if key == 'dns_csv' and not value and isinstance(site.get('dns'), list):
                value = ' | '.join(str(v) for v in site['dns'])
            ws.cell(row=row_idx, column=col_idx, value='' if value is None else value)

    for col_idx in range(1, len(EXPORT_COLUMNS) + 1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = 18 if col_idx < 7 else 40

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, filename
